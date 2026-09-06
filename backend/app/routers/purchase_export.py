from fastapi import APIRouter, Depends, HTTPException
from fastapi.responses import StreamingResponse
from sqlalchemy.orm import Session
from sqlalchemy import extract
from datetime import datetime
import pandas as pd
import io
from openpyxl.styles import Font, Alignment
from .. import models
from ..database import get_db
from ..config import get_settings

router = APIRouter(prefix="/purchase-export", tags=["purchase-export"])


def build_purchase_export_data(purchases):
    """Build export data for purchases. ONLY purchase-side data is included."""
    data = []
    for purchase in purchases:
        for item in purchase.items:
            data.append({
                "Purchase Invoice No": purchase.bill_number,
                "Date": purchase.invoice_date.strftime("%Y-%m-%d") if purchase.invoice_date else "",
                "Supplier Name": purchase.supplier.name if purchase.supplier else "Unknown",
                "Supplier GSTIN": (purchase.supplier.gstin or "") if purchase.supplier else "",
                "Product Name": item.description,
                "HSN": item.hsn_code or "",
                "Quantity": item.quantity,
                "Unit": item.unit or "",
                "Rate": item.rate,
                "Taxable Amount": round(item.taxable_amount, 2),
                "GST %": item.gst_percentage,
                "CGST": round(item.cgst, 2),
                "SGST": round(item.sgst, 2),
                "Total": round(item.amount, 2),
                "Amount Paid": round(purchase.amount_paid, 2),
                "Balance Due": round(purchase.balance_due, 2),
                "Payment Status": purchase.payment_status
            })
    return data


def write_purchase_excel_with_header(df, output, sub_title):
    """Write purchase DataFrame to Excel with formatted header and a numeric TOTAL row.

    Layout (Excel 1-indexed):
      Row 1: business name (merged)
      Row 2: subtitle (merged)
      Row 3: spacer
      Row 4: column headers (pandas startrow=3)
      Row 5..: one row per purchase line item
      Last data row + 1: TOTAL row (numeric sums, bold)
    """
    with pd.ExcelWriter(output, engine='openpyxl') as writer:
        df.to_excel(writer, index=False, sheet_name='Purchases', startrow=3)

        worksheet = writer.sheets['Purchases']
        settings = get_settings()

        max_col = len(df.columns)

        worksheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=max_col)
        worksheet.merge_cells(start_row=2, start_column=1, end_row=2, end_column=max_col)

        cell1 = worksheet.cell(row=1, column=1)
        cell1.value = settings.business_name
        cell1.font = Font(size=14, bold=True)
        cell1.alignment = Alignment(horizontal='center')

        cell2 = worksheet.cell(row=2, column=1)
        cell2.value = sub_title
        cell2.font = Font(size=12, bold=True)
        cell2.alignment = Alignment(horizontal='center')

        # Data starts at Excel row 5 (pandas startrow=3 -> header row 4, data row 5+)
        start_row = 5
        data_rows = len(df)

        if data_rows > 0:
            last_data_row = start_row + data_rows - 1
            total_row = last_data_row + 1

            # Column mapping (Excel columns, 1-indexed):
            # 1=PurchaseInvoice, 2=Date, 3=Supplier, 4=SupplierGSTIN, 5=Product,
            # 6=HSN, 7=Quantity, 8=Unit, 9=Rate, 10=Taxable, 11=GST%,
            # 12=CGST, 13=SGST, 14=Total, 15=AmountPaid, 16=BalanceDue, 17=Status
            col_taxable = 10
            col_cgst = 12
            col_sgst = 13
            col_total = 14
            col_amount_paid = 15
            col_balance_due = 16
            col_product = 5

            # "TOTAL" label spanning product columns only (columns 5-9)
            worksheet.cell(row=total_row, column=col_product).value = "TOTAL"
            worksheet.cell(row=total_row, column=col_product).font = Font(bold=True)
            worksheet.merge_cells(
                start_row=total_row, start_column=col_product,
                end_row=total_row, end_column=col_product + 4
            )

            # Numeric totals computed in Python (guaranteed to display).
            taxable_sum = round(df['Taxable Amount'].sum(), 2)
            cgst_sum = round(df['CGST'].sum(), 2)
            sgst_sum = round(df['SGST'].sum(), 2)
            total_sum = round(df['Total'].sum(), 2)
            amount_paid_sum = round(df['Amount Paid'].sum(), 2)
            balance_due_sum = round(df['Balance Due'].sum(), 2)

            for col, val in [
                (col_taxable, taxable_sum),
                (col_cgst, cgst_sum),
                (col_sgst, sgst_sum),
                (col_total, total_sum),
                (col_amount_paid, amount_paid_sum),
                (col_balance_due, balance_due_sum),
            ]:
                cell = worksheet.cell(row=total_row, column=col)
                cell.value = val
                cell.font = Font(bold=True)


@router.get("/monthly")
def export_monthly_purchases(month: int, year: int, db: Session = Depends(get_db)):
    if month < 1 or month > 12:
        raise HTTPException(status_code=400, detail="Invalid month")

    purchases = db.query(models.Purchase).filter(
        extract('month', models.Purchase.invoice_date) == month,
        extract('year', models.Purchase.invoice_date) == year
    ).all()

    if not purchases:
        raise HTTPException(status_code=404, detail="No purchases found for the specified month")

    data = build_purchase_export_data(purchases)
    df = pd.DataFrame(data)

    output = io.BytesIO()

    month_names = ['', 'January', 'February', 'March', 'April', 'May', 'June',
                   'July', 'August', 'September', 'October', 'November', 'December']
    write_purchase_excel_with_header(df, output, f"{month_names[month]} {year} Purchases")

    output.seek(0)

    filename = f"purchases_{year}_{month:02d}.xlsx"

    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={filename}"}
    )


@router.get("/all")
def export_all_purchases(db: Session = Depends(get_db)):
    purchases = db.query(models.Purchase).order_by(models.Purchase.invoice_date.desc()).all()

    if not purchases:
        raise HTTPException(status_code=404, detail="No purchases found")

    data = build_purchase_export_data(purchases)
    df = pd.DataFrame(data)

    output = io.BytesIO()
    write_purchase_excel_with_header(df, output, "All Purchases")

    output.seek(0)

    filename = f"all_purchases_{datetime.now().strftime('%Y%m%d')}.xlsx"

    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={filename}"}
    )